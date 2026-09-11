from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent
p=ROOT/'manuscript.ko.md'
s=p.read_text().replace('EdiVal-Agent: An Object-Centric Framework for Automated Evaluation of Image Editing.','EdiVal-Agent: An Object-Centric Framework for Automated, Fine-Grained Evaluation of Multi-Turn Editing.')
p.write_text(s)
d=Document(); sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=d.styles[name];st.font.name='Noto Sans CJK KR';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR');st.font.color.rgb=RGBColor.from_string('172333')
st=d.styles['Normal'];st.font.size=Pt(10);st.paragraph_format.line_spacing=1.18;st.paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(20);d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 2'].font.size=Pt(11.5)
for name in ['Heading 1','Heading 2']:d.styles[name].paragraph_format.space_before=Pt(12)
pf=sec.footer.paragraphs[0];pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');pf._p.append(f)
lines=s.splitlines();i=0;tn=0
captions=['실험 묶음과 평가 범위','P04 정책별 호출 수와 경로 평균 거리','후속 인물별 얼굴 LPIPS','후속 사람·AI 평가 분포','로컬 모델 최종 얼굴 지표','로컬 영역·위치 보정 민감도']
while i<len(lines):
 line=lines[i]
 if not line:i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   r=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.fullmatch('[-: ]+',v) for v in r):rows.append(r)
   i+=1
  cap=d.add_paragraph(f'표 {tn+1}. {captions[tn]}','Caption');cap.paragraph_format.keep_with_next=True;tn+=1
  t=d.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
  for j,row in enumerate(rows):
   cells=t.add_row().cells
   pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
   if j==0:pr.append(OxmlElement('w:tblHeader'))
   for c,txt in zip(cells,row):
    c.text=txt
    for para in c.paragraphs:
     para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4);para.paragraph_format.line_spacing=1.05
     for run in para.runs:run.font.size=Pt(8 if len(row)>4 else 9);run.bold=j==0
  d.add_paragraph().paragraph_format.space_after=Pt(1)
  continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  d.add_picture(str(ROOT/m[2]),width=Cm(16.4));d.paragraphs[-1].paragraph_format.keep_with_next=True
  d.add_paragraph(m[1],'Caption')
 elif line.startswith('# '):d.add_paragraph(line[2:].replace('원본 기반 재생성 비교','\n원본 기반 재생성 비교'),'Title')
 elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 1')
 else:d.add_paragraph(line)
 i+=1
d.core_properties.title=s.splitlines()[0][2:];d.core_properties.subject='Exploratory empirical study';d.core_properties.author=''
for style in d.styles:
 if style.type == 1:
  style.font.name='Noto Sans CJK KR'
  rf=style._element.get_or_add_rPr().rFonts
  for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Noto Sans CJK KR')
  for key in list(rf.attrib):
   if 'Theme' in key:del rf.attrib[key]
for el in list(d._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for st in d.styles:
 for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
d.save(ROOT/'manuscript.ko.docx')
print('DOCX built',tn,'tables')
