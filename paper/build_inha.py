from pathlib import Path
import re,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent
s=(R/'manuscript.ko.md').read_text().split('## 부록 A.')[0].strip()
abstract='We compare sequential image editing with regeneration from the original image and accumulated requirements. Experiments include six synthetic identities, 98 follow-up outputs from three identities, automatic image metrics and separate AI ratings. Batch editing produced lower mean facial LPIPS in the tested follow-up comparisons. A local FLUX.2 Klein 4B comparison yielded mean LPIPS of 0.091313 for batch and 0.124455 for sequential editing, while raw skin MAE reversed the ranking in some settings. A web editor implemented the workflow. Recorded checks covered real generation, state changes, version restoration and 35 known request regression cases. The findings characterize the dependence of input-policy effects on measurement regions and support an editor that maintains accumulated editing requirements.'
auth=json.loads((R/'authors.json').read_text()) if (R/'authors.json').exists() else {'korean':'________________','english':'________________','advisor_korean':'________________','advisor_english':'________________'}
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10);st.font.color.rgb=RGBColor(0,0,0);rf=st._element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK KR')
 for k in list(rf.attrib):
  if 'Theme' in k:del rf.attrib[k]
 for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
n=d.styles['Normal'].paragraph_format;n.line_spacing=Pt(14);n.space_after=Pt(5);n.first_line_indent=Cm(.3);n.widow_control=True
for name in ['Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Noto Sans CJK KR';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR');st.font.bold=True;st.font.size=Pt(11 if name=='Heading 1' else 10);st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
d.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['Title','Subtitle']:d.styles[name].paragraph_format.first_line_indent=Pt(0)
d.styles['Title'].font.size=Pt(17);d.styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR');d.styles['Title'].font.bold=True
d.styles['Subtitle'].font.italic=False;d.styles['Subtitle'].font.size=Pt(14);d.styles['Subtitle'].font.bold=True

def center(text,style=None):
 p=d.add_paragraph(text,style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);return p

def columns(num):
 sec=d.add_section(WD_SECTION_START.CONTINUOUS);cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),str(num));cols.set(qn('w:space'),'500');return sec

center(s.splitlines()[0][2:].replace('원본 기반 재생성 비교','\n원본 기반 재생성 비교'),'Title').paragraph_format.line_spacing=Pt(24)
center(s.splitlines()[2],'Subtitle').paragraph_format.line_spacing=Pt(18)
center(auth['korean']).paragraph_format.space_before=Pt(12)
center('('+auth['english']+')')
center('지도교수: '+auth['advisor_korean']).paragraph_format.space_before=Pt(6)
center('('+auth['advisor_english']+')').paragraph_format.space_after=Pt(14)
ko=s.split('## 초록\n')[1].split('\n\n주요어:')[0].strip()
for label,txt in [('요약: ',ko),('Abstract: ',abstract),('Keywords: ','Sequential image editing, Facial preservation, Original-referenced regeneration, Perceptual similarity, Exploratory evaluation')]:
 p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.add_run(label).bold=True;p.add_run(txt)
columns(2)
lines=s.split('## 1. 서론')[1];lines='## 1. 서론'+lines
lines=lines.splitlines();i=0;tn=0;roman=['I','II','III','IV','V','VI','VII']
caps=['실험과 평가 범위','P04 정책별 호출·LPIPS','인물별 평균 얼굴 LPIPS','후속 AI 평가','로컬 최종 얼굴 지표','피부 영역·위치 보정 민감도']
while i<len(lines):
 line=lines[i]
 if not line:i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   rr=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.fullmatch('[-: ]+',v) for v in rr):rows.append(rr)
   i+=1
  tn+=1;wide=tn in [1,6]
  if wide:columns(1)
  p=d.add_paragraph(f'표 {tn}. {caps[tn-1]}','Caption');p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if tn==2:rows[0]=['정책','호출','고정','보정']
  if tn==3:rows[0]=['인물','일괄','순차','반복']
  if tn==5:rows[0]=['시드·방법','MAE','SSIM','LPIPS']
  t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=False
  width=17 if wide else 8.06
  for col in t.columns:col.width=Cm(width/len(rows[0]))
  for j,row in enumerate(rows):
   cells=t.add_row().cells;pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
   if j==0:pr.append(OxmlElement('w:tblHeader'))
   for c,txt in zip(cells,row):
    c.text=txt
    for p in c.paragraphs:
     p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
     for run in p.runs:run.font.size=Pt(8);run.bold=j==0
    if j==0 or j==len(rows)-1:
     borders=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');borders.append(e);c._tc.get_or_add_tcPr().append(borders)
  if wide:columns(2)
  continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line);columns(1);d.add_picture(str(R/m[2]),width=Cm(16.8));d.paragraphs[-1].paragraph_format.keep_with_next=True;d.paragraphs[-1].paragraph_format.line_spacing=1;d.add_paragraph(m[1],'Caption');columns(2)
 elif line.startswith('### '):d.add_paragraph(re.sub(r'^\d+\.(\d+) ',r'\1. ',line[4:]),'Heading 2')
 elif line.startswith('## '):
  text=line[3:];text=re.sub(r'^(\d+)\.',lambda m:roman[int(m[1])-1]+'.',text);d.add_paragraph(text,'Heading 1')
 else:
  p=d.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if line.startswith('['):p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True
 i+=1
for sec in d.sections:
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0)
 if not f._p.find(qn('w:fldSimple')) is not None:
  e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),'PAGE');f._p.append(e)
d.core_properties.title=s.splitlines()[0][2:];d.core_properties.author='' if '_' in auth['korean'] else auth['korean']
d.save(R/'manuscript.inha.docx')
(R/'manuscript.inha.md').write_text(s.replace('## 초록','저자: '+auth['korean']+' ('+auth['english']+')\n\n지도교수: '+auth['advisor_korean']+'\n\n## 초록').replace('## 1. 서론','## Abstract\n\n'+abstract+'\n\n## 1. 서론')+'\n')
print('Built Inha reference format')
