"""Builder regressions; run with the bundled document Python runtime."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None


@unittest.skipIf(Document is None, 'Use the bundled document runtime with python-docx')
class BuilderTests(unittest.TestCase):
    def test_private_source_metadata_and_small_table_pagination(self):
        with tempfile.TemporaryDirectory(prefix='inha-builder-test-') as folder:
            root = Path(folder)
            source, output = root / 'source', root / 'output'
            source.mkdir()
            (source / 'manuscript.ko.md').write_text(
                '# 검증 제목\n\nValidation Title\n\n## 초록\n\n검증 요약.\n\n주요어: 검증\n\n'
                '## 1. 서론\n\n검증 본문.\n\n'
                '| 방법 | 정확도 |\n| --- | --- |\n| 기준 | 1/2 |\n| 제안 | 2/2 |\n\n'
                '## 참고문헌\n\n[1] Example reference.\n')
            (source / 'abstract.en.txt').write_text('Validation abstract.\n')
            (source / 'layout.json').write_text(json.dumps({
                'keywords': 'Validation keyword', 'table_captions': ['검증 표']}))
            subprocess.run([sys.executable, str(Path(__file__).with_name('build_inha.py')),
                            '--source-dir', str(source), '--output-dir', str(output)],
                           check=True, capture_output=True)
            doc = Document(output / 'manuscript.inha.docx')
            paragraphs = [p.text for p in doc.paragraphs]
            self.assertIn('검증 제목', paragraphs)
            self.assertIn('표 1. 검증 표', paragraphs)
            self.assertIn('Keywords: Validation keyword', paragraphs)
            self.assertFalse((source / 'manuscript.inha.docx').exists())
            table = doc.tables[0]
            self.assertEqual(len(table.rows), 3)
            for row in table.rows[:-1]:
                self.assertTrue(row.cells[0].paragraphs[0].paragraph_format.keep_with_next)
            self.assertFalse(table.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next)
            self.assertEqual(table.rows[1].cells[1].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            reference = next(p for p in doc.paragraphs if p.text.startswith('[1]'))
            self.assertEqual(reference.alignment, WD_ALIGN_PARAGRAPH.LEFT)


if __name__ == '__main__':
    unittest.main()
